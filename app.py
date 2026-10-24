"""Transaction sharing web app.

Users register with a unique user ID and password. A user uploads an Excel
sheet of bank transactions; each row is delivered to the user named in its
"Recipient" column. Recipients can view their transactions and download them
as a PDF.
"""

import io
import os
import re
import secrets
import sqlite3
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from functools import wraps

from flask import (
    Flask,
    abort,
    current_app,
    flash,
    g,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from werkzeug.security import check_password_hash, generate_password_hash

USER_ID_RE = re.compile(r"^[A-Za-z0-9_.-]{3,32}$")
MIN_PASSWORD_LEN = 6
MAX_ROWS = 5000

# Accepted header spellings (lower-cased) for each column of the upload sheet.
COLUMN_ALIASES = {
    "recipient": {"recipient", "recipient id", "recipient user id", "user id", "user", "to", "to user"},
    "date": {"date", "transaction date", "txn date"},
    "description": {"description", "details", "narration", "particulars", "remarks"},
    "type": {"type", "transaction type", "txn type", "dr/cr", "debit/credit"},
    "amount": {"amount", "amt", "value"},
    "reference": {"reference", "ref", "reference no", "reference number", "cheque no", "txn id"},
}
REQUIRED_COLUMNS = ("recipient", "amount")

SCHEMA = """
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id TEXT NOT NULL UNIQUE COLLATE NOCASE,
    password_hash TEXT NOT NULL,
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS uploads (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    uploader_id INTEGER NOT NULL REFERENCES users(id),
    filename TEXT NOT NULL,
    row_count INTEGER NOT NULL,
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS transactions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    upload_id INTEGER NOT NULL REFERENCES uploads(id),
    sender_id INTEGER NOT NULL REFERENCES users(id),
    recipient_id INTEGER NOT NULL REFERENCES users(id),
    txn_date TEXT,
    description TEXT,
    txn_type TEXT,
    amount_cents INTEGER NOT NULL,
    reference TEXT,
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_txn_recipient ON transactions(recipient_id);
CREATE INDEX IF NOT EXISTS idx_txn_sender ON transactions(sender_id);
"""


def create_app(config=None):
    app = Flask(__name__, instance_relative_config=True)
    app.config.update(
        SECRET_KEY=os.environ.get("SECRET_KEY") or secrets.token_hex(32),
        DATABASE=os.environ.get("DATABASE") or os.path.join(app.instance_path, "transactions.db"),
        MAX_CONTENT_LENGTH=5 * 1024 * 1024,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
    )
    if config:
        app.config.update(config)
    os.makedirs(os.path.dirname(os.path.abspath(app.config["DATABASE"])), exist_ok=True)

    app.teardown_appcontext(close_db)
    with app.app_context():
        get_db().executescript(SCHEMA)

    app.jinja_env.filters["money"] = format_money
    register_routes(app)
    return app


# --------------------------------------------------------------------------- db


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(current_app.config["DATABASE"])
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


def close_db(_exc=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def now_iso():
    return datetime.now().isoformat(timespec="seconds")


# ---------------------------------------------------------------------- helpers


def format_money(cents):
    sign = "-" if cents < 0 else ""
    cents = abs(cents)
    return f"{sign}{cents // 100:,}.{cents % 100:02d}"


def current_user():
    uid = session.get("uid")
    if uid is None:
        return None
    if "user" not in g:
        g.user = get_db().execute("SELECT id, user_id FROM users WHERE id = ?", (uid,)).fetchone()
    return g.user


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if current_user() is None:
            session.clear()
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)

    return wrapped


def csrf_token():
    if "csrf" not in session:
        session["csrf"] = secrets.token_urlsafe(32)
    return session["csrf"]


def check_csrf():
    token = request.form.get("csrf_token", "")
    if not token or not secrets.compare_digest(token, session.get("csrf", "")):
        abort(400, "Invalid or missing form token. Reload the page and try again.")


def safe_next(target):
    if target and target.startswith("/") and not target.startswith("//"):
        return target
    return url_for("dashboard")


# ------------------------------------------------------------------ excel parse


class UploadError(Exception):
    pass


def parse_amount(value):
    if value is None or value == "":
        raise ValueError("amount is empty")
    if isinstance(value, (int, float, Decimal)):
        amount = Decimal(str(value))
    else:
        text = str(value).strip().replace(",", "")
        text = re.sub(r"^[^\d\-+.(]+", "", text)  # drop currency symbols / prefixes
        negative = text.startswith("(") and text.endswith(")")
        text = text.strip("()")
        try:
            amount = Decimal(text)
        except InvalidOperation:
            raise ValueError(f"'{value}' is not a valid amount") from None
        if negative:
            amount = -amount
    return int((amount * 100).quantize(Decimal("1")))


def parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d.%m.%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text  # keep unrecognised formats verbatim rather than rejecting the row


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_workbook(stream):
    """Return a list of dicts (one per non-empty row) from the first sheet."""
    try:
        wb = load_workbook(stream, read_only=True, data_only=True)
    except Exception:
        raise UploadError("Could not read the file. Please upload a valid .xlsx Excel workbook.") from None

    try:
        rows = wb.active.iter_rows(values_only=True)
        header = None
        for row in rows:
            if any(cell_text(c) for c in row):
                header = row
                break
        if header is None:
            raise UploadError("The spreadsheet is empty.")

        columns = {}
        for idx, name in enumerate(header):
            key = cell_text(name).lower()
            for field, aliases in COLUMN_ALIASES.items():
                if key in aliases and field not in columns:
                    columns[field] = idx
        missing = [c for c in REQUIRED_COLUMNS if c not in columns]
        if missing:
            raise UploadError(
                "Missing required column(s): " + ", ".join(m.title() for m in missing)
                + ". Download the template to see the expected layout."
            )

        records = []
        for line_no, row in enumerate(rows, start=2):
            if not any(cell_text(c) for c in row):
                continue
            get = lambda f: row[columns[f]] if f in columns and columns[f] < len(row) else None  # noqa: E731
            records.append({"line": line_no, **{f: get(f) for f in COLUMN_ALIASES}})
            if len(records) > MAX_ROWS:
                raise UploadError(f"Too many rows; the limit is {MAX_ROWS} per upload.")
        if not records:
            raise UploadError("The spreadsheet has a header row but no transactions.")
        return records
    finally:
        wb.close()


# ----------------------------------------------------------------------- pdf


def build_pdf(title, subtitle, rows, counterpart_label):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=15 * mm, title=title,
    )
    styles = getSampleStyleSheet()
    cell = styles["BodyText"]
    cell.fontSize = 9
    cell.leading = 11

    data = [["#", "Date", counterpart_label, "Description", "Type", "Reference", "Amount"]]
    total = 0
    for i, r in enumerate(rows, start=1):
        total += r["amount_cents"]
        data.append([
            str(i),
            r["txn_date"] or "",
            r["counterpart"],
            Paragraph(escape_xml(r["description"] or ""), cell),
            r["txn_type"] or "",
            Paragraph(escape_xml(r["reference"] or ""), cell),
            format_money(r["amount_cents"]),
        ])
    data.append(["", "", "", "", "", "Total", format_money(total)])

    table = Table(
        data, repeatRows=1,
        colWidths=[12 * mm, 25 * mm, 35 * mm, 95 * mm, 22 * mm, 40 * mm, 30 * mm],
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f4e79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (-1, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f2f6fa")]),
        ("GRID", (0, 0), (-1, -2), 0.25, colors.HexColor("#c9d3dd")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("LINEABOVE", (0, -1), (-1, -1), 1, colors.HexColor("#1f4e79")),
    ]))

    story = [
        Paragraph(escape_xml(title), styles["Title"]),
        Paragraph(escape_xml(subtitle), styles["Normal"]),
        Spacer(1, 6 * mm),
    ]
    story.append(table if rows else Paragraph("No transactions.", styles["Normal"]))
    doc.build(story)
    buf.seek(0)
    return buf


def escape_xml(text):
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# --------------------------------------------------------------------- routes


TXN_SELECT = """
SELECT t.*, s.user_id AS sender, r.user_id AS recipient
FROM transactions t
JOIN users s ON s.id = t.sender_id
JOIN users r ON r.id = t.recipient_id
"""


def register_routes(app):
    @app.context_processor
    def inject():
        return {"current_user": current_user(), "csrf_token": csrf_token}

    @app.route("/")
    def index():
        return redirect(url_for("dashboard" if current_user() else "login"))

    @app.route("/register", methods=["GET", "POST"])
    def register():
        if current_user():
            return redirect(url_for("dashboard"))
        if request.method == "POST":
            check_csrf()
            user_id = request.form.get("user_id", "").strip()
            password = request.form.get("password", "")
            confirm = request.form.get("confirm", "")
            error = None
            if not USER_ID_RE.match(user_id):
                error = "User ID must be 3-32 characters: letters, numbers, '.', '_' or '-'."
            elif len(password) < MIN_PASSWORD_LEN:
                error = f"Password must be at least {MIN_PASSWORD_LEN} characters."
            elif password != confirm:
                error = "Passwords do not match."
            else:
                db = get_db()
                try:
                    cur = db.execute(
                        "INSERT INTO users (user_id, password_hash, created_at) VALUES (?, ?, ?)",
                        (user_id, generate_password_hash(password), now_iso()),
                    )
                    db.commit()
                except sqlite3.IntegrityError:
                    error = f"The user ID '{user_id}' is already taken. Please choose another."
                else:
                    session.clear()
                    session["uid"] = cur.lastrowid
                    flash(f"Welcome, {user_id}! Your account has been created.", "success")
                    return redirect(url_for("dashboard"))
            flash(error, "error")
            return render_template("register.html", user_id=user_id), 400
        return render_template("register.html", user_id="")

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if current_user():
            return redirect(url_for("dashboard"))
        if request.method == "POST":
            check_csrf()
            user_id = request.form.get("user_id", "").strip()
            password = request.form.get("password", "")
            user = get_db().execute(
                "SELECT id, password_hash FROM users WHERE user_id = ?", (user_id,)
            ).fetchone()
            if user and check_password_hash(user["password_hash"], password):
                session.clear()
                session["uid"] = user["id"]
                return redirect(safe_next(request.args.get("next")))
            flash("Invalid user ID or password.", "error")
            return render_template("login.html", user_id=user_id), 401
        return render_template("login.html", user_id="")

    @app.route("/logout", methods=["POST"])
    def logout():
        check_csrf()
        session.clear()
        flash("You have been logged out.", "success")
        return redirect(url_for("login"))

    @app.route("/dashboard")
    @login_required
    def dashboard():
        db = get_db()
        uid = current_user()["id"]
        received = db.execute(
            TXN_SELECT + " WHERE t.recipient_id = ? ORDER BY t.txn_date DESC, t.id DESC", (uid,)
        ).fetchall()
        uploads = db.execute(
            "SELECT * FROM uploads WHERE uploader_id = ? ORDER BY id DESC", (uid,)
        ).fetchall()
        senders = db.execute(
            """SELECT DISTINCT u.user_id FROM transactions t JOIN users u ON u.id = t.sender_id
               WHERE t.recipient_id = ? ORDER BY u.user_id""", (uid,)
        ).fetchall()
        return render_template(
            "dashboard.html",
            received=received,
            received_total=sum(r["amount_cents"] for r in received),
            uploads=uploads,
            senders=[s["user_id"] for s in senders],
        )

    @app.route("/upload", methods=["GET", "POST"])
    @login_required
    def upload():
        if request.method == "GET":
            return render_template("upload.html", result=None)

        check_csrf()
        file = request.files.get("file")
        if not file or not file.filename:
            flash("Please choose an Excel file to upload.", "error")
            return render_template("upload.html", result=None), 400
        if not file.filename.lower().endswith((".xlsx", ".xlsm")):
            flash("Only Excel .xlsx files are supported.", "error")
            return render_template("upload.html", result=None), 400

        try:
            records = parse_workbook(io.BytesIO(file.read()))
        except UploadError as exc:
            flash(str(exc), "error")
            return render_template("upload.html", result=None), 400

        db = get_db()
        me = current_user()
        users = db.execute("SELECT id, user_id FROM users").fetchall()
        user_ids = {u["user_id"].lower(): u["id"] for u in users}
        display_names = {u["id"]: u["user_id"] for u in users}

        valid, skipped = [], []
        for rec in records:
            recipient = cell_text(rec["recipient"])
            if not recipient:
                skipped.append((rec["line"], "Recipient is empty"))
                continue
            recipient_pk = user_ids.get(recipient.lower())
            if recipient_pk is None:
                skipped.append((rec["line"], f"No user with ID '{recipient}'"))
                continue
            try:
                cents = parse_amount(rec["amount"])
            except ValueError as exc:
                skipped.append((rec["line"], str(exc).capitalize()))
                continue
            valid.append((
                recipient_pk,
                parse_date(rec["date"]),
                cell_text(rec["description"])[:500],
                cell_text(rec["type"])[:30],
                cents,
                cell_text(rec["reference"])[:100],
            ))

        delivered = {}
        if valid:
            now = now_iso()
            cur = db.execute(
                "INSERT INTO uploads (uploader_id, filename, row_count, created_at) VALUES (?, ?, ?, ?)",
                (me["id"], os.path.basename(file.filename)[:200], len(valid), now),
            )
            upload_id = cur.lastrowid
            db.executemany(
                """INSERT INTO transactions (upload_id, sender_id, recipient_id, txn_date,
                   description, txn_type, amount_cents, reference, created_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                [(upload_id, me["id"], *v, now) for v in valid],
            )
            db.commit()
            for v in valid:
                name = display_names[v[0]]
                delivered[name] = delivered.get(name, 0) + 1
            flash(f"Uploaded {len(valid)} transaction(s) to {len(delivered)} user(s).", "success")
        else:
            flash("No transactions were uploaded. Fix the problems below and try again.", "error")

        return render_template(
            "upload.html",
            result={"delivered": sorted(delivered.items()), "skipped": skipped, "total": len(records)},
        )

    @app.route("/template.xlsx")
    @login_required
    def template():
        wb = Workbook()
        ws = wb.active
        ws.title = "Transactions"
        ws.append(["Recipient", "Date", "Description", "Type", "Amount", "Reference"])
        ws.append(["alice", date.today(), "Invoice #1001 payment", "Credit", 2500.00, "TXN-0001"])
        ws.append(["bob", date.today(), "Office supplies", "Debit", 149.99, "TXN-0002"])
        for col, width in zip("ABCDEF", (16, 12, 36, 10, 12, 16)):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
            row[0].number_format = "yyyy-mm-dd"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(
            buf, as_attachment=True, download_name="transactions_template.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.route("/transactions.pdf")
    @login_required
    def received_pdf():
        """PDF of transactions the current user received, optionally filtered by sender."""
        me = current_user()
        sql = TXN_SELECT + " WHERE t.recipient_id = ?"
        params = [me["id"]]
        sender = request.args.get("sender", "").strip()
        if sender:
            sql += " AND s.user_id = ?"
            params.append(sender)
        sql += " ORDER BY t.txn_date, t.id"
        rows = [
            {**dict(r), "counterpart": r["sender"]}
            for r in get_db().execute(sql, params).fetchall()
        ]
        subtitle = f"Recipient: {me['user_id']}"
        if sender:
            subtitle += f" | From: {sender}"
        subtitle += f" | Generated: {datetime.now():%Y-%m-%d %H:%M}"
        pdf = build_pdf("Transaction Statement", subtitle, rows, "From")
        name = f"transactions_{me['user_id']}" + (f"_from_{sender}" if sender else "") + ".pdf"
        return send_file(pdf, as_attachment=True, download_name=name, mimetype="application/pdf")

    @app.route("/uploads/<int:upload_id>")
    @login_required
    def upload_detail(upload_id):
        db = get_db()
        up = db.execute(
            "SELECT * FROM uploads WHERE id = ? AND uploader_id = ?", (upload_id, current_user()["id"])
        ).fetchone()
        if up is None:
            abort(404)
        rows = db.execute(TXN_SELECT + " WHERE t.upload_id = ? ORDER BY t.id", (upload_id,)).fetchall()
        return render_template("upload_detail.html", upload=up, rows=rows,
                               total=sum(r["amount_cents"] for r in rows))

    @app.route("/uploads/<int:upload_id>.pdf")
    @login_required
    def upload_pdf(upload_id):
        db = get_db()
        me = current_user()
        up = db.execute(
            "SELECT * FROM uploads WHERE id = ? AND uploader_id = ?", (upload_id, me["id"])
        ).fetchone()
        if up is None:
            abort(404)
        rows = [
            {**dict(r), "counterpart": r["recipient"]}
            for r in db.execute(TXN_SELECT + " WHERE t.upload_id = ? ORDER BY t.id", (upload_id,))
        ]
        subtitle = f"Uploaded by: {me['user_id']} | File: {up['filename']} | {up['created_at']}"
        pdf = build_pdf("Uploaded Transactions", subtitle, rows, "To")
        return send_file(pdf, as_attachment=True, download_name=f"upload_{upload_id}.pdf",
                         mimetype="application/pdf")

    @app.errorhandler(413)
    def too_large(_e):
        flash("File is too large (limit 5 MB).", "error")
        return redirect(url_for("upload"))


if __name__ == "__main__":
    create_app().run(debug=os.environ.get("FLASK_DEBUG") == "1", host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
