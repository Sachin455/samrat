"""Admin portal: manage user IDs, upload Excel files, view and edit every transaction."""

import io
import os
import re
from datetime import date, datetime
from functools import wraps

from flask import (
    Blueprint,
    abort,
    flash,
    g,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook
from werkzeug.security import check_password_hash

from core import (
    MIN_PASSWORD_LEN,
    PERSONAL_FIELDS,
    PROFILE_FIELDS,
    TXN_ORDER,
    TXN_SELECT,
    UploadError,
    build_pdf,
    cell_text,
    cents_to_input,
    create_users,
    get_db,
    get_profile,
    get_setting,
    hash_password,
    normalize_type,
    normalize_user_id,
    now_iso,
    parse_amount,
    parse_date,
    parse_workbook,
    record_amount_and_type,
    safe_next,
    save_profile,
    set_setting,
    txn_summary,
    validate_new_password,
    validate_profile,
)

bp = Blueprint("admin", __name__, url_prefix="/admin")
PAGE_SIZE = 50


def current_admin():
    aid = session.get("admin_id")
    if aid is None:
        return None
    if "admin" not in g:
        g.admin = get_db().execute(
            "SELECT id, username, must_change_password FROM admins WHERE id = ?", (aid,)
        ).fetchone()
    return g.admin


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        admin = current_admin()
        if admin is None:
            session.pop("admin_id", None)
            return redirect(url_for("admin.login", next=request.path))
        if admin["must_change_password"] and request.endpoint != "admin.change_password":
            return redirect(url_for("admin.change_password"))
        return view(*args, **kwargs)

    return wrapped


def get_user_or_404(user_id):
    user = get_db().execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone()
    if user is None:
        abort(404)
    return user


def get_txn_or_404(txn_id):
    txn = get_db().execute(f"{TXN_SELECT} WHERE t.id = ?", (txn_id,)).fetchone()
    if txn is None:
        abort(404)
    return txn


def parse_ids(text):
    """Split pasted text into IDs; return (valid_ids, invalid_tokens)."""
    valid, invalid = [], []
    for token in re.split(r"[\s,;]+", text or ""):
        if not token:
            continue
        uid = normalize_user_id(token)
        (valid if uid else invalid).append(uid or token)
    return valid, invalid


def txn_form_values(form):
    """Validate a transaction form. Returns (values, error)."""
    user_id = normalize_user_id(form.get("user_id", ""))
    if not user_id:
        return None, "User ID must be exactly 9 digits."
    user = get_db().execute("SELECT id FROM users WHERE user_id = ?", (user_id,)).fetchone()
    if user is None:
        return None, f"No user with ID {user_id}. Add the user first."
    try:
        cents = parse_amount(form.get("amount", ""))
    except ValueError as exc:
        return None, f"Amount: {exc}."
    return {
        "user_pk": user["id"],
        "txn_date": parse_date(form.get("txn_date", "").strip()),
        "description": form.get("description", "").strip()[:500],
        "txn_type": normalize_type(form.get("txn_type", "")),
        "amount_cents": cents,
        "reference": form.get("reference", "").strip()[:100],
    }, None


# ------------------------------------------------------------------ auth


@bp.route("/login", methods=["GET", "POST"])
def login():
    if current_admin():
        return redirect(url_for("admin.dashboard"))
    if request.method == "POST":
        username = request.form.get("user_id", "").strip()
        password = request.form.get("password", "")
        admin = get_db().execute(
            "SELECT id, password_hash, must_change_password FROM admins WHERE username = ?", (username,)
        ).fetchone()
        if admin and check_password_hash(admin["password_hash"], password):
            session["admin_id"] = admin["id"]
            if admin["must_change_password"]:
                flash("Please change the admin password before continuing.", "success")
                return redirect(url_for("admin.change_password"))
            return redirect(safe_next(request.args.get("next"), url_for("admin.dashboard")))
        flash("Invalid admin username or password.", "error")
        return render_template("login.html", portal="admin", user_id=username), 401
    return render_template("login.html", portal="admin", user_id="")


@bp.route("/logout", methods=["POST"])
def logout():
    session.pop("admin_id", None)
    flash("You have been logged out of the admin portal.", "success")
    return redirect(url_for("admin.login"))


@bp.route("/change-password", methods=["GET", "POST"])
@admin_required
def change_password():
    admin = current_admin()
    if request.method == "POST":
        db = get_db()
        stored = db.execute("SELECT password_hash FROM admins WHERE id = ?", (admin["id"],)).fetchone()
        current = request.form.get("current_password", "")
        new = request.form.get("new_password", "")
        error = validate_new_password(
            check_password_hash(stored["password_hash"], current), new, request.form.get("confirm", ""), current
        )
        if error:
            flash(error, "error")
            return render_template("change_password.html", portal="admin",
                                   forced=admin["must_change_password"]), 400
        db.execute(
            "UPDATE admins SET password_hash = ?, must_change_password = 0 WHERE id = ?",
            (hash_password(new), admin["id"]),
        )
        db.commit()
        flash("Admin password changed.", "success")
        return redirect(url_for("admin.dashboard"))
    return render_template("change_password.html", portal="admin", forced=admin["must_change_password"])


# ------------------------------------------------------------- dashboard


@bp.route("/")
@admin_required
def dashboard():
    db = get_db()
    return render_template(
        "admin/dashboard.html",
        user_count=db.execute("SELECT COUNT(*) FROM users").fetchone()[0],
        summary=txn_summary("1 = 1", ()),
        latest=db.execute(f"{TXN_SELECT} ORDER BY t.id DESC LIMIT 10").fetchall(),
        uploads=db.execute("SELECT * FROM uploads ORDER BY id DESC LIMIT 5").fetchall(),
    )


# ----------------------------------------------------------------- users


@bp.route("/users", methods=["GET", "POST"])
@admin_required
def users():
    db = get_db()
    if request.method == "POST":
        action = request.form.get("action")
        if action == "add":
            valid, invalid = parse_ids(request.form.get("user_ids", ""))
            if invalid:
                flash("Not 9-digit IDs (ignored): " + ", ".join(invalid[:20]), "error")
            if valid:
                created, existing = create_users(valid)
                db.commit()
                if created:
                    flash(f"Added {len(created)} user(s): " + ", ".join(created[:20])
                          + ("…" if len(created) > 20 else ""), "success")
                if existing:
                    flash("Already existed: " + ", ".join(existing[:20]), "error")
            elif not invalid:
                flash("Enter at least one 9-digit user ID.", "error")
        elif action == "default_password":
            new_default = request.form.get("default_password", "")
            if len(new_default) < MIN_PASSWORD_LEN:
                flash(f"Default password must be at least {MIN_PASSWORD_LEN} characters.", "error")
            else:
                set_setting("default_password", new_default)
                db.commit()
                flash("Default first-time password updated. It applies to users added or reset from now on.",
                      "success")
        return redirect(url_for("admin.users", q=request.args.get("q", "")))

    q = request.args.get("q", "").strip()
    rows = db.execute(
        """SELECT u.*, p.full_name, COUNT(t.id) AS txn_count, MAX(t.txn_date) AS last_txn
           FROM users u
           LEFT JOIN profiles p ON p.user_pk = u.id
           LEFT JOIN transactions t ON t.user_pk = u.id
           WHERE u.user_id LIKE ? OR p.full_name LIKE ? GROUP BY u.id ORDER BY u.user_id""",
        (f"%{q}%", f"%{q}%"),
    ).fetchall()
    return render_template("admin/users.html", users=rows, q=q, default_password=get_setting("default_password"))


@bp.route("/users/<user_id>")
@admin_required
def user_detail(user_id):
    user = get_user_or_404(user_id)
    rows = get_db().execute(f"{TXN_SELECT} WHERE t.user_pk = ? {TXN_ORDER}", (user["id"],)).fetchall()
    return render_template(
        "admin/user_detail.html", user=user, rows=rows, summary=txn_summary("t.user_pk = ?", (user["id"],)),
        today=date.today().isoformat(), profile=get_profile(user["id"]), fields=PROFILE_FIELDS,
        personal=get_profile(user["id"], "personal"), personal_fields=PERSONAL_FIELDS,
    )


@bp.route("/users/<user_id>/profile", methods=["GET", "POST"])
@admin_required
def edit_profile(user_id):
    user = get_user_or_404(user_id)
    page = {"portal": "admin", "fields": PROFILE_FIELDS, "heading": f"Organization profile of user {user_id}",
            "intro": "", "first_time": False, "cancel_url": url_for("admin.user_detail", user_id=user_id)}
    if request.method == "POST":
        values, errors = validate_profile(request.form)
        if errors:
            flash("Please correct the highlighted fields.", "error")
            return render_template("profile_form.html", values=values, errors=errors, **page), 400
        save_profile(user["id"], values)
        get_db().commit()
        flash(f"Profile of {user_id} saved.", "success")
        return redirect(url_for("admin.user_detail", user_id=user_id))
    existing = get_profile(user["id"])
    return render_template("profile_form.html", values=dict(existing) if existing else {}, errors={}, **page)


@bp.route("/users/<user_id>/transactions.pdf")
@admin_required
def user_pdf(user_id):
    user = get_user_or_404(user_id)
    rows = get_db().execute(f"{TXN_SELECT} WHERE t.user_pk = ? {TXN_ORDER}", (user["id"],)).fetchall()
    profile = get_profile(user["id"])
    name = f"{profile['full_name']} | " if profile else ""
    subtitle = f"{name}User ID: {user['user_id']} | Generated: {datetime.now():%Y-%m-%d %H:%M}"
    return send_file(build_pdf("Transaction Statement", subtitle, rows), as_attachment=True,
                     download_name=f"statement_{user['user_id']}.pdf", mimetype="application/pdf")


@bp.route("/users/<user_id>/reset-password", methods=["POST"])
@admin_required
def reset_password(user_id):
    user = get_user_or_404(user_id)
    db = get_db()
    db.execute(
        "UPDATE users SET password_hash = ?, must_change_password = 1 WHERE id = ?",
        (hash_password(get_setting("default_password")), user["id"]),
    )
    db.commit()
    flash(f"Password for {user_id} reset to the default first-time password.", "success")
    return redirect(url_for("admin.user_detail", user_id=user_id))


@bp.route("/users/<user_id>/delete", methods=["POST"])
@admin_required
def delete_user(user_id):
    user = get_user_or_404(user_id)
    db = get_db()
    db.execute("DELETE FROM users WHERE id = ?", (user["id"],))
    db.commit()
    flash(f"User {user_id} and all their transactions were deleted.", "success")
    return redirect(url_for("admin.users"))


# ---------------------------------------------------------- transactions


@bp.route("/transactions")
@admin_required
def transactions():
    q = re.sub(r"\D", "", request.args.get("q", ""))
    page = max(request.args.get("page", 1, type=int), 1)
    where, params = "u.user_id LIKE ?", [f"%{q}%"]
    db = get_db()
    total = db.execute(
        f"SELECT COUNT(*) FROM transactions t JOIN users u ON u.id = t.user_pk WHERE {where}", params
    ).fetchone()[0]
    rows = db.execute(
        f"{TXN_SELECT} WHERE {where} {TXN_ORDER} LIMIT ? OFFSET ?", params + [PAGE_SIZE, (page - 1) * PAGE_SIZE]
    ).fetchall()
    pages = max((total + PAGE_SIZE - 1) // PAGE_SIZE, 1)
    return render_template("admin/transactions.html", rows=rows, q=q, page=page, pages=pages, total=total)


@bp.route("/transactions/new", methods=["POST"])
@admin_required
def new_transaction():
    values, error = txn_form_values(request.form)
    back = url_for("admin.user_detail", user_id=request.form.get("user_id", "")) \
        if request.form.get("from_user") else url_for("admin.transactions")
    if error:
        flash(error, "error")
        return redirect(back)
    db = get_db()
    db.execute(
        """INSERT INTO transactions (user_pk, txn_date, description, txn_type, amount_cents, reference, created_at)
           VALUES (:user_pk, :txn_date, :description, :txn_type, :amount_cents, :reference, :now)""",
        {**values, "now": now_iso()},
    )
    db.commit()
    flash("Transaction added.", "success")
    return redirect(back)


@bp.route("/transactions/<int:txn_id>/edit", methods=["GET", "POST"])
@admin_required
def edit_transaction(txn_id):
    txn = get_txn_or_404(txn_id)
    if request.method == "POST":
        values, error = txn_form_values(request.form)
        if error:
            flash(error, "error")
            return render_template("admin/transaction_form.html", txn=txn, form=request.form), 400
        db = get_db()
        db.execute(
            """UPDATE transactions SET user_pk = :user_pk, txn_date = :txn_date, description = :description,
               txn_type = :txn_type, amount_cents = :amount_cents, reference = :reference, updated_at = :now
               WHERE id = :id""",
            {**values, "now": now_iso(), "id": txn_id},
        )
        db.commit()
        flash("Transaction updated.", "success")
        return redirect(safe_next(request.args.get("next"), url_for("admin.transactions")))
    form = {
        "user_id": txn["user_id"], "txn_date": txn["txn_date"] or "", "description": txn["description"],
        "txn_type": txn["txn_type"], "amount": cents_to_input(txn["amount_cents"]), "reference": txn["reference"],
    }
    return render_template("admin/transaction_form.html", txn=txn, form=form)


@bp.route("/transactions/<int:txn_id>/delete", methods=["POST"])
@admin_required
def delete_transaction(txn_id):
    get_txn_or_404(txn_id)
    db = get_db()
    db.execute("DELETE FROM transactions WHERE id = ?", (txn_id,))
    db.commit()
    flash("Transaction deleted.", "success")
    return redirect(safe_next(request.args.get("next"), url_for("admin.transactions")))


# --------------------------------------------------------------- uploads


@bp.route("/upload", methods=["GET", "POST"])
@admin_required
def upload():
    if request.method == "GET":
        return render_template("admin/upload.html", result=None, form={"create_missing": "1"})

    form = request.form
    target_raw = form.get("target_user_id", "").strip()
    target = normalize_user_id(target_raw) if target_raw else None
    create_missing = bool(form.get("create_missing"))

    def fail(message):
        flash(message, "error")
        return render_template("admin/upload.html", result=None, form=form), 400

    if target_raw and not target:
        return fail("User ID must be exactly 9 digits.")
    file = request.files.get("file")
    if not file or not file.filename:
        return fail("Please choose an Excel file to upload.")
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        return fail("Only Excel .xlsx files are supported.")
    try:
        records = parse_workbook(io.BytesIO(file.read()), need_user_column=target is None)
    except UploadError as exc:
        return fail(str(exc))

    db = get_db()
    valid, skipped = [], []
    for rec in records:
        if target:
            uid = target
        else:
            uid = normalize_user_id(rec["user"])
            if uid is None:
                shown = cell_text(rec["user"])
                skipped.append((rec["line"], f"'{shown}' is not a 9-digit user ID" if shown else "User ID is empty"))
                continue
        try:
            cents, txn_type = record_amount_and_type(rec)
        except ValueError as exc:
            skipped.append((rec["line"], str(exc).capitalize()))
            continue
        valid.append((rec["line"], uid, parse_date(rec["date"]), cell_text(rec["description"])[:500],
                      txn_type, cents, cell_text(rec["reference"])[:100]))

    wanted = {v[1] for v in valid}
    if wanted:
        placeholders = ",".join("?" * len(wanted))
        known = {r["user_id"]: r["id"] for r in db.execute(
            f"SELECT id, user_id FROM users WHERE user_id IN ({placeholders})", list(wanted))}
    else:
        known = {}
    unknown = sorted(wanted - known.keys())
    created = []
    if unknown and create_missing:
        created, _ = create_users(unknown)
    if created:
        known.update({r["user_id"]: r["id"] for r in db.execute(
            f"SELECT id, user_id FROM users WHERE user_id IN ({','.join('?' * len(created))})", created)})
    rows = []
    for line, uid, *rest in valid:
        if uid not in known:
            skipped.append((line, f"No user with ID {uid} (tick “create missing users” to add it)"))
            continue
        rows.append((known[uid], *rest))
    skipped.sort()

    delivered = {}
    if rows:
        now = now_iso()
        cur = db.execute(
            "INSERT INTO uploads (filename, target_user_id, row_count, created_at) VALUES (?, ?, ?, ?)",
            (os.path.basename(file.filename)[:200], target, len(rows), now),
        )
        db.executemany(
            """INSERT INTO transactions (user_pk, upload_id, txn_date, description, txn_type,
               amount_cents, reference, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            [(r[0], cur.lastrowid, *r[1:], now) for r in rows],
        )
        db.commit()
        pk_to_id = {pk: uid for uid, pk in known.items()}
        for r in rows:
            delivered[pk_to_id[r[0]]] = delivered.get(pk_to_id[r[0]], 0) + 1
        flash(f"Uploaded {len(rows)} transaction(s) to {len(delivered)} user(s).", "success")
    else:
        db.rollback()  # don't keep users created for an upload that stored nothing
        created = []
        flash("No transactions were uploaded. Fix the problems below and try again.", "error")

    return render_template(
        "admin/upload.html", form=form,
        result={"delivered": sorted(delivered.items()), "skipped": skipped, "total": len(records),
                "created": created},
    )


@bp.route("/uploads/<int:upload_id>")
@admin_required
def upload_detail(upload_id):
    db = get_db()
    up = db.execute("SELECT * FROM uploads WHERE id = ?", (upload_id,)).fetchone()
    if up is None:
        abort(404)
    rows = db.execute(f"{TXN_SELECT} WHERE t.upload_id = ? ORDER BY t.id", (upload_id,)).fetchall()
    return render_template("admin/upload_detail.html", upload=up, rows=rows)


@bp.route("/uploads/<int:upload_id>/delete", methods=["POST"])
@admin_required
def delete_upload(upload_id):
    db = get_db()
    if db.execute("SELECT 1 FROM uploads WHERE id = ?", (upload_id,)).fetchone() is None:
        abort(404)
    n = db.execute("DELETE FROM transactions WHERE upload_id = ?", (upload_id,)).rowcount
    db.execute("DELETE FROM uploads WHERE id = ?", (upload_id,))
    db.commit()
    flash(f"Upload deleted along with its {n} transaction(s).", "success")
    return redirect(url_for("admin.dashboard"))


@bp.route("/template.xlsx")
@admin_required
def template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["User ID", "Date", "Description", "Type", "Amount", "Reference"])
    ws.append(["100000001", date.today(), "Salary October", "Credit", 2500.00, "TXN-0001"])
    ws.append(["100000002", date.today(), "ATM withdrawal", "Debit", 149.99, "TXN-0002"])
    for col, width in zip("ABCDEF", (14, 12, 36, 10, 12, 16)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, max_col=2):
        row[0].number_format = "@"  # keep user IDs as text so leading zeros survive
        row[1].number_format = "yyyy-mm-dd"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="transactions_template.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
