"""Shared pieces: database, security helpers, Excel parsing and PDF output."""

import io
import os
import re
import secrets
import sqlite3
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from flask import abort, current_app, g, request, session
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from werkzeug.security import generate_password_hash

USER_ID_RE = re.compile(r"^\d{9}$")
MIN_PASSWORD_LEN = 6
# Werkzeug defaults to scrypt, which Python builds linked against LibreSSL
# (e.g. macOS's stock Python 3.9) lack; PBKDF2 is available everywhere.
PASSWORD_HASH_METHOD = "pbkdf2:sha256"
MAX_ROWS = 5000
LATEST_COUNT = 5

DEFAULT_USER_PASSWORD = "Welcome@123"
DEFAULT_ADMIN_USERNAME = "admin"
DEFAULT_ADMIN_PASSWORD = "admin123"

SCHEMA = """
CREATE TABLE IF NOT EXISTS settings (
    key TEXT PRIMARY KEY,
    value TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS admins (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    username TEXT NOT NULL UNIQUE COLLATE NOCASE,
    password_hash TEXT NOT NULL,
    must_change_password INTEGER NOT NULL DEFAULT 1
);
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id TEXT NOT NULL UNIQUE,
    password_hash TEXT NOT NULL,
    must_change_password INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL,
    last_login TEXT
);
CREATE TABLE IF NOT EXISTS uploads (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    filename TEXT NOT NULL,
    target_user_id TEXT,
    row_count INTEGER NOT NULL,
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS transactions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    user_pk INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
    upload_id INTEGER REFERENCES uploads(id) ON DELETE SET NULL,
    txn_date TEXT,
    description TEXT NOT NULL DEFAULT '',
    txn_type TEXT NOT NULL DEFAULT '',
    amount_cents INTEGER NOT NULL,
    reference TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    updated_at TEXT
);
CREATE TABLE IF NOT EXISTS profiles (
    user_pk INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,
    full_name TEXT NOT NULL,
    address TEXT NOT NULL,
    established_date TEXT NOT NULL,
    head_name TEXT NOT NULL,
    email TEXT NOT NULL,
    phone TEXT NOT NULL,
    running_classes TEXT NOT NULL,
    chairperson TEXT NOT NULL,
    accountant TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS personal_profiles (
    user_pk INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,
    full_name TEXT NOT NULL,
    designation TEXT NOT NULL DEFAULT '',
    email TEXT NOT NULL DEFAULT '',
    phone TEXT NOT NULL DEFAULT '',
    address TEXT NOT NULL DEFAULT '',
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_txn_user ON transactions(user_pk, txn_date);
CREATE INDEX IF NOT EXISTS idx_txn_upload ON transactions(upload_id);
"""

# Newest first; rows without a date go last.
TXN_ORDER = "ORDER BY t.txn_date IS NULL, t.txn_date DESC, t.id DESC"
DEBIT_SQL = "lower(t.txn_type) LIKE 'd%' OR lower(t.txn_type) LIKE 'withdraw%'"
TXN_SELECT = "SELECT t.*, u.user_id FROM transactions t JOIN users u ON u.id = t.user_pk"


# --------------------------------------------------------------------------- db


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(current_app.config["DATABASE"])
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


def close_db(_exc=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = get_db()
    db.executescript(SCHEMA)
    if db.execute("SELECT 1 FROM admins").fetchone() is None:
        db.execute(
            "INSERT INTO admins (username, password_hash) VALUES (?, ?)",
            (
                os.environ.get("ADMIN_USERNAME") or DEFAULT_ADMIN_USERNAME,
                hash_password(os.environ.get("ADMIN_PASSWORD") or DEFAULT_ADMIN_PASSWORD),
            ),
        )
    db.execute(
        "INSERT OR IGNORE INTO settings (key, value) VALUES ('default_password', ?)",
        (os.environ.get("DEFAULT_USER_PASSWORD") or DEFAULT_USER_PASSWORD,),
    )
    db.commit()


def get_setting(key):
    row = get_db().execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
    return row["value"] if row else None


def set_setting(key, value):
    get_db().execute(
        "INSERT INTO settings (key, value) VALUES (?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
        (key, value),
    )


def now_iso():
    return datetime.now().isoformat(timespec="seconds")


def hash_password(password):
    return generate_password_hash(password, method=PASSWORD_HASH_METHOD)


def create_users(user_ids):
    """Create accounts with the current default password. Returns (created, existing)."""
    db = get_db()
    password_hash = hash_password(get_setting("default_password"))
    created, existing = [], []
    for uid in dict.fromkeys(user_ids):  # de-duplicate, keep order
        try:
            db.execute(
                "INSERT INTO users (user_id, password_hash, created_at) VALUES (?, ?, ?)",
                (uid, password_hash, now_iso()),
            )
            created.append(uid)
        except sqlite3.IntegrityError:
            existing.append(uid)
    return created, existing


# (column, label, input type, max length, required) for the two profile forms.
PROFILE_FIELDS = [  # organization profile, filled in on first login
    ("full_name", "Full name of the organization", "text", 200, True),
    ("address", "Address", "textarea", 500, True),
    ("established_date", "Establishment date", "date", 10, True),
    ("head_name", "Head of the organization", "text", 120, True),
    ("email", "Organization email", "email", 120, True),
    ("phone", "Organization contact number", "tel", 30, True),
    ("running_classes", "Running classes", "textarea", 500, True),
    ("chairperson", "Chairperson of the organization", "text", 120, True),
    ("accountant", "Accountant of the organization", "text", 120, True),
]
PERSONAL_FIELDS = [  # the account holder's own details, editable any time
    ("full_name", "Full name", "text", 120, True),
    ("designation", "Designation / position", "text", 120, False),
    ("email", "Personal email", "email", 120, False),
    ("phone", "Mobile number", "tel", 30, False),
    ("address", "Personal address", "textarea", 500, False),
]
PROFILE_TABLES = {"organization": ("profiles", PROFILE_FIELDS), "personal": ("personal_profiles", PERSONAL_FIELDS)}
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PHONE_RE = re.compile(r"^\+?[\d\s()-]+$")


def get_profile(user_pk, kind="organization"):
    table, _ = PROFILE_TABLES[kind]
    return get_db().execute(f"SELECT * FROM {table} WHERE user_pk = ?", (user_pk,)).fetchone()


def validate_profile(form, kind="organization"):
    """Return (values, errors) where errors maps field name -> message."""
    _, fields = PROFILE_TABLES[kind]
    values, errors = {}, {}
    for key, label, input_type, max_len, required in fields:
        value = form.get(key, "").strip()
        values[key] = value
        if not value:
            if required:
                errors[key] = f"{label} is required."
            continue
        if len(value) > max_len:
            errors[key] = f"{label} must be at most {max_len} characters."
        elif input_type == "email" and not EMAIL_RE.match(value):
            errors[key] = "Enter a valid email address, e.g. office@school.edu.np."
        elif input_type == "tel" and (not PHONE_RE.match(value) or not 7 <= len(re.sub(r"\D", "", value)) <= 15):
            errors[key] = "Enter a valid phone number (7 to 15 digits)."
        elif input_type == "date":
            try:
                day = datetime.strptime(value, "%Y-%m-%d").date()
            except ValueError:
                errors[key] = f"Enter the {label.lower()} as a date."
            else:
                if day > date.today():
                    errors[key] = f"{label} can't be in the future."
    return values, errors


def save_profile(user_pk, values, kind="organization"):
    table, fields = PROFILE_TABLES[kind]
    columns = [key for key, *_ in fields]
    get_db().execute(
        f"""INSERT INTO {table} (user_pk, {", ".join(columns)}, updated_at)
            VALUES (?, {", ".join("?" * len(columns))}, ?)
            ON CONFLICT(user_pk) DO UPDATE SET
            {", ".join(f"{c} = excluded.{c}" for c in columns)}, updated_at = excluded.updated_at""",
        (user_pk, *(values[c] for c in columns), now_iso()),
    )


def txn_summary(where, params):
    row = get_db().execute(
        f"""SELECT COUNT(*) AS n,
                   COALESCE(SUM(CASE WHEN {DEBIT_SQL} THEN 0 ELSE amount_cents END), 0) AS credits,
                   COALESCE(SUM(CASE WHEN {DEBIT_SQL} THEN amount_cents ELSE 0 END), 0) AS debits
            FROM transactions t WHERE {where}""",
        params,
    ).fetchone()
    return {"count": row["n"], "credits": row["credits"], "debits": row["debits"]}


# ---------------------------------------------------------------------- security


def csrf_token():
    if "csrf" not in session:
        session["csrf"] = secrets.token_urlsafe(32)
    return session["csrf"]


def check_csrf():
    token = request.form.get("csrf_token", "")
    if not token or not secrets.compare_digest(token, session.get("csrf", "")):
        abort(400, "Invalid or missing form token. Reload the page and try again.")


def safe_next(target, fallback):
    if target and target.startswith("/") and not target.startswith("//"):
        return target
    return fallback


def validate_new_password(current_ok, new, confirm, old_password):
    if not current_ok:
        return "Your current password is incorrect."
    if len(new) < MIN_PASSWORD_LEN:
        return f"New password must be at least {MIN_PASSWORD_LEN} characters."
    if new != confirm:
        return "New passwords do not match."
    if new == old_password:
        return "New password must be different from the current one."
    return None


# ------------------------------------------------------------------- formatting


def format_money(cents):
    sign = "-" if cents < 0 else ""
    cents = abs(cents)
    return f"{sign}{cents // 100:,}.{cents % 100:02d}"


def cents_to_input(cents):
    return format_money(cents).replace(",", "")


def is_debit(txn_type):
    t = (txn_type or "").strip().lower()
    return t.startswith("d") or t.startswith("withdraw")


def normalize_type(value):
    text = cell_text(value)
    low = text.lower()
    if low in ("d", "dr", "debit", "withdrawal", "withdraw"):
        return "Debit"
    if low in ("c", "cr", "credit", "deposit"):
        return "Credit"
    return text[:30]


# ------------------------------------------------------------------- parsing


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime):
        value = value.date()
    return str(value).strip()


def normalize_user_id(value):
    """Return a 9-digit user ID string, or None if the value isn't one.

    Excel stores IDs typed as numbers without leading zeros, so numeric cells
    are zero-padded back to 9 digits.
    """
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        text = str(int(value)).zfill(9)
    else:
        text = re.sub(r"[\s-]", "", str(value))
    return text if USER_ID_RE.match(text) else None


def parse_amount(value):
    if value is None or value == "":
        raise ValueError("amount is empty")
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        amount = Decimal(str(value))
    else:
        text = str(value).strip().replace(",", "")
        text = re.sub(r"^[^\d\-+.(]+", "", text)  # drop currency symbols / prefixes
        negative = text.startswith("(") and text.endswith(")")
        text = text.strip("()")
        try:
            amount = Decimal(text)
        except InvalidOperation:
            raise ValueError(f"'{value}' is not a valid amount") from None
        if negative:
            amount = -amount
    if not amount.is_finite():
        raise ValueError(f"'{value}' is not a valid amount")
    return int((amount * 100).quantize(Decimal("1")))


def parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d.%m.%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text  # keep unrecognised formats verbatim rather than rejecting the row


# Accepted header spellings (lower-cased) for each column of the upload sheet.
COLUMN_ALIASES = {
    "user": {"user id", "userid", "user", "customer id", "account id", "id", "recipient", "recipient id", "to"},
    "date": {"date", "transaction date", "txn date", "value date"},
    "description": {"description", "details", "narration", "particulars", "remarks"},
    "type": {"type", "transaction type", "txn type", "dr/cr", "debit/credit"},
    "amount": {"amount", "amt", "value"},
    "debit": {"debit", "withdrawal", "dr", "debit amount", "withdrawal amount"},
    "credit": {"credit", "deposit", "cr", "credit amount", "deposit amount"},
    "reference": {"reference", "ref", "reference no", "reference number", "cheque no", "txn id"},
}


class UploadError(Exception):
    pass


def parse_workbook(stream, need_user_column):
    """Return a list of dicts (one per non-empty row) from the first sheet."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(stream, read_only=True, data_only=True)
    except Exception:
        raise UploadError("Could not read the file. Please upload a valid .xlsx Excel workbook.") from None

    try:
        rows = wb.active.iter_rows(values_only=True)
        header = next((r for r in rows if any(cell_text(c) for c in r)), None)
        if header is None:
            raise UploadError("The spreadsheet is empty.")

        columns = {}
        for idx, name in enumerate(header):
            key = cell_text(name).lower()
            for field, aliases in COLUMN_ALIASES.items():
                if key in aliases and field not in columns:
                    columns[field] = idx
        missing = []
        if need_user_column and "user" not in columns:
            missing.append("User ID")
        if not ({"amount", "debit", "credit"} & columns.keys()):
            missing.append("Amount (or Debit/Credit)")
        if missing:
            raise UploadError(
                "Missing required column(s): " + ", ".join(missing)
                + ". Download the template to see the expected layout."
            )

        records = []
        for line_no, row in enumerate(rows, start=2):
            if not any(cell_text(c) for c in row):
                continue
            rec = {"line": line_no}
            for field in COLUMN_ALIASES:
                idx = columns.get(field)
                rec[field] = row[idx] if idx is not None and idx < len(row) else None
            records.append(rec)
            if len(records) > MAX_ROWS:
                raise UploadError(f"Too many rows; the limit is {MAX_ROWS} per upload.")
        if not records:
            raise UploadError("The spreadsheet has a header row but no transactions.")
        return records
    finally:
        wb.close()


def record_amount_and_type(rec):
    """Work out (amount_cents, type) from either an Amount column or Debit/Credit columns."""
    if cell_text(rec["amount"]):
        return parse_amount(rec["amount"]), normalize_type(rec["type"])
    if cell_text(rec["debit"]):
        return abs(parse_amount(rec["debit"])), "Debit"
    if cell_text(rec["credit"]):
        return abs(parse_amount(rec["credit"])), "Credit"
    raise ValueError("amount is empty")


# ----------------------------------------------------------------------- pdf


def escape_xml(text):
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(title, subtitle, rows, include_user=False):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=15 * mm, title=title,
    )
    styles = getSampleStyleSheet()
    cell = styles["BodyText"]
    cell.fontSize = 9
    cell.leading = 11

    header = ["#", "Date"] + (["User ID"] if include_user else []) + ["Description", "Type", "Reference", "Amount"]
    widths = [12, 25] + ([28] if include_user else []) + [(112 if include_user else 140), 22, 40, 28]
    data = [header]
    credits = debits = 0
    for i, r in enumerate(rows, start=1):
        if is_debit(r["txn_type"]):
            debits += r["amount_cents"]
        else:
            credits += r["amount_cents"]
        data.append(
            [str(i), r["txn_date"] or ""]
            + ([r["user_id"]] if include_user else [])
            + [
                Paragraph(escape_xml(r["description"] or ""), cell),
                r["txn_type"] or "",
                Paragraph(escape_xml(r["reference"] or ""), cell),
                format_money(r["amount_cents"]),
            ]
        )

    table = Table(data, repeatRows=1, colWidths=[w * mm for w in widths])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f4e79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (-1, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f6fa")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c9d3dd")),
    ]))

    totals = Table(
        [["Total credits", format_money(credits)],
         ["Total debits", format_money(debits)],
         ["Net (credits - debits)", format_money(credits - debits)]],
        colWidths=[50 * mm, 30 * mm], hAlign="RIGHT",
    )
    totals.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("LINEABOVE", (0, -1), (-1, -1), 1, colors.HexColor("#1f4e79")),
    ]))

    story = [
        Paragraph(escape_xml(title), styles["Title"]),
        Paragraph(escape_xml(subtitle), styles["Normal"]),
        Spacer(1, 6 * mm),
    ]
    if rows:
        story += [table, Spacer(1, 4 * mm), totals]
    else:
        story.append(Paragraph("No transactions.", styles["Normal"]))
    doc.build(story)
    buf.seek(0)
    return buf
